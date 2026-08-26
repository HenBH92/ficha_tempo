"""Leitura das listas auxiliares do modelo e gravação de linhas na planilha de saída.
A planilha de saída só existe no arquivo que o usuário escolher salvar (sem cache local)."""
import shutil
import warnings
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

from caminhos import pasta_dados, pasta_recursos

CAMINHO_MODELO = pasta_recursos() / "aquivo_modelo.xlsx"
CAMINHO_LOG_ADVWIN = pasta_dados() / "log_advwin.xlsx"  # histórico do que já foi lançado no AdvWin

ABAS_LISTAS = ("Advogado", "Cliente")

# O modelo usa validação de dados em extensão x14 (listas grandes demais para o
# formato clássico do Excel) que o openpyxl não sabe reescrever e descarta ao
# salvar. Os valores das colunas continuam corretos (o app já valida antes de
# gravar); só o dropdown nativo do Excel na planilha de saída se perde.
warnings.filterwarnings("ignore", message=".*Data Validation extension.*")


def carregar_listas(caminho_modelo: Path = CAMINHO_MODELO) -> dict[str, list[str]]:
    wb = openpyxl.load_workbook(caminho_modelo, read_only=True, data_only=True)
    listas = {}
    for aba in ABAS_LISTAS:
        ws = wb[aba]
        listas[aba] = [c[0].value for c in ws.iter_rows(min_col=1, max_col=1) if c[0].value]
    wb.close()
    return listas


def garantir_saida(caminho_modelo: Path = CAMINHO_MODELO, *, caminho_saida_arquivo: Path) -> Path:
    """Se o arquivo escolhido pelo usuário ainda não existe, semeia com o modelo (abas/colunas/listas)."""
    if not caminho_saida_arquivo.exists():
        shutil.copy(caminho_modelo, caminho_saida_arquivo)
    return caminho_saida_arquivo


def _proxima_linha_vazia(ws) -> int:
    linha = 3
    while ws.cell(linha, 2).value not in (None, ""):
        linha += 1
    return linha


def inserir_linhas(lista_dados: list[dict], caminho_saida_arquivo: Path) -> None:
    """Grava várias linhas de uma vez (abre/salva o arquivo uma única vez)."""
    caminho_saida_arquivo = garantir_saida(caminho_saida_arquivo=caminho_saida_arquivo)
    wb = openpyxl.load_workbook(caminho_saida_arquivo)
    ws = wb["Ficha-tempo"]
    linha = _proxima_linha_vazia(ws)

    for dados in lista_dados:
        ws.cell(linha, 1, "Ficha Tempo")
        ws.cell(linha, 2, dados["data"])
        ws.cell(linha, 3, dados["advogado"])
        ws.cell(linha, 4, "-")
        ws.cell(linha, 5, dados["cliente"])
        ws.cell(linha, 6, str(dados["pasta"]))
        ws.cell(linha, 7, dados["descricao"])
        ws.cell(linha, 8, dados["horas"])
        ws.cell(linha, 9, dados["horas_cobraveis"])
        ws.cell(linha, 16, "MATRIZ")
        linha += 1

    wb.save(caminho_saida_arquivo)


def inserir_linha(dados: dict, caminho_saida_arquivo: Path) -> None:
    inserir_linhas([dados], caminho_saida_arquivo)


def parse_horas_cobraveis(texto: str) -> timedelta | None:
    """Aceita "H:MM" (ex.: "1:05") ou só minutos (ex.: "35")."""
    texto = texto.strip()
    if not texto:
        return None
    try:
        if ":" in texto:
            h, m = texto.split(":", 1)
            return timedelta(hours=int(h), minutes=int(m))
        return timedelta(minutes=float(texto.replace(",", ".")))
    except ValueError:
        return None


def _demo():
    import os
    import tempfile

    tmp = Path(tempfile.gettempdir()) / "ficha_tempo_teste.xlsx"
    if tmp.exists():
        os.remove(tmp)

    listas = carregar_listas()
    assert len(listas["Advogado"]) > 0
    assert len(listas["Cliente"]) > 1000

    garantir_saida(caminho_saida_arquivo=tmp)
    assert tmp.exists()

    agora = datetime.now()
    linha_base = {
        "data": agora.date(),
        "advogado": listas["Advogado"][0],
        "cliente": listas["Cliente"][0],
        "pasta": "123VT-CIV-0001.01",
        "descricao": "teste",
        "horas": timedelta(minutes=37),
        "horas_cobraveis": timedelta(minutes=40),
    }
    inserir_linha(linha_base, caminho_saida_arquivo=tmp)
    inserir_linhas([linha_base, linha_base], caminho_saida_arquivo=tmp)  # grava 2 de uma vez

    wb = openpyxl.load_workbook(tmp)
    ws = wb["Ficha-tempo"]
    assert ws.cell(3, 3).value == listas["Advogado"][0]
    assert ws.cell(4, 3).value == listas["Advogado"][0]
    assert ws.cell(5, 3).value == listas["Advogado"][0]
    assert ws.cell(6, 3).value in (None, "")  # não sobrou linha extra
    os.remove(tmp)
    print("planilha.py OK")


if __name__ == "__main__":
    _demo()
