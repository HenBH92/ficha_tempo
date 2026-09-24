"""Lançamento retroativo em lote no AdvWin, lendo linhas já preenchidas numa cópia da
planilha de modelo (aquivo_modelo.xlsx, aba "Ficha-tempo" - mesmas colunas que planilha.py
já escreve). Cada linha processada com sucesso é marcada na própria planilha (coluna 17),
pra permitir reabrir o mesmo arquivo sem lançar a mesma linha duas vezes - o AdvWin não
permite desfazer um lançamento."""
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl

ABA = "Ficha-tempo"
COL_DATA, COL_ADVOGADO, COL_CLIENTE, COL_PASTA, COL_DESCRICAO, COL_HORAS = 2, 3, 5, 6, 7, 8
COL_STATUS = 17  # uma coluna depois da última usada pelo modelo (16 = Unidade)
CABECALHO_STATUS = "Status AdvWin"
PRIMEIRA_LINHA = 3
PREFIXO_SUCESSO = "Lançado em"


@dataclass
class LinhaRetroativa:
    linha: int
    pasta: str
    data: str
    descricao: str
    horas_texto: str
    advogado: str
    cliente: str


def _texto_data(valor) -> str:
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def _texto_horas(valor) -> str:
    """"Horas" (coluna 8) pode voltar do Excel como time, timedelta ou fração de dia,
    dependendo do formato da célula - normaliza tudo pra "h:mm"."""
    if valor is None:
        return ""
    if isinstance(valor, timedelta):
        minutos = round(valor.total_seconds() / 60)
    elif isinstance(valor, time):
        minutos = valor.hour * 60 + valor.minute
    elif isinstance(valor, (int, float)):
        minutos = round(valor * 24 * 60)
    else:
        return str(valor).strip()
    return f"{minutos // 60}:{minutos % 60:02d}"


def ler_linhas_pendentes(caminho: Path) -> list[LinhaRetroativa]:
    """Linhas com Data preenchida (mesmo critério de fim-de-dados de planilha.py) que ainda
    não têm marca de sucesso na coluna de status - linhas com erro numa tentativa anterior
    entram de novo, pra dar chance de corrigir e reprocessar."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[ABA]
    linhas = []
    linha = PRIMEIRA_LINHA
    while ws.cell(linha, COL_DATA).value not in (None, ""):
        status = ws.cell(linha, COL_STATUS).value
        if not (isinstance(status, str) and status.startswith(PREFIXO_SUCESSO)):
            linhas.append(LinhaRetroativa(
                linha=linha,
                pasta=str(ws.cell(linha, COL_PASTA).value or "").strip(),
                data=_texto_data(ws.cell(linha, COL_DATA).value),
                descricao=str(ws.cell(linha, COL_DESCRICAO).value or "").strip(),
                horas_texto=_texto_horas(ws.cell(linha, COL_HORAS).value),
                advogado=str(ws.cell(linha, COL_ADVOGADO).value or "").strip(),
                cliente=str(ws.cell(linha, COL_CLIENTE).value or "").strip(),
            ))
        linha += 1
    wb.close()
    return linhas


def marcar_linha(caminho: Path, linha: int, texto_status: str) -> None:
    """Reabre e salva a planilha pra cada linha - garante que o progresso fica gravado em
    disco mesmo se o app fechar no meio do lote."""
    wb = openpyxl.load_workbook(caminho)
    ws = wb[ABA]
    if ws.cell(2, COL_STATUS).value != CABECALHO_STATUS:
        ws.cell(2, COL_STATUS, CABECALHO_STATUS)
    ws.cell(linha, COL_STATUS, texto_status)
    wb.save(caminho)
    wb.close()


def _demo():
    import os
    import shutil
    import tempfile

    import planilha

    tmp = Path(tempfile.gettempdir()) / "ficha_tempo_retroativo_teste.xlsx"
    if tmp.exists():
        os.remove(tmp)
    shutil.copy(planilha.CAMINHO_MODELO, tmp)

    linha_base = {
        "data": date(2026, 8, 10),
        "advogado": "Fulano",
        "cliente": "Cliente Teste",
        "pasta": "123VT-CIV-0001.01",
        "descricao": "teste retroativo",
        "horas": timedelta(minutes=30),
        "horas_cobraveis": timedelta(minutes=35),
    }
    planilha.inserir_linhas([linha_base, linha_base], caminho_saida_arquivo=tmp)

    pendentes = ler_linhas_pendentes(tmp)
    assert len(pendentes) == 2, pendentes
    assert pendentes[0].pasta == "123VT-CIV-0001.01"
    assert pendentes[0].data == "10/08/2026"
    assert pendentes[0].horas_texto == "0:30"  # coluna "Horas" (8), não a cobrável (9)
    assert pendentes[0].advogado == "Fulano"

    marcar_linha(tmp, pendentes[0].linha, "Lançado em 24/08/2026 10:00")
    marcar_linha(tmp, pendentes[1].linha, "Erro: pasta não encontrada")

    restantes = ler_linhas_pendentes(tmp)
    assert len(restantes) == 1  # a marcada como erro volta a aparecer; a de sucesso não
    assert restantes[0].linha == pendentes[1].linha

    os.remove(tmp)
    print("retroativo.py OK")


if __name__ == "__main__":
    _demo()
